import os
import json
import time
import gspread
import argparse
import pandas as pd
from collections import Counter
from google.cloud import storage
from google.auth import default as google_auth_default
from google.auth import impersonated_credentials
from openpyxl.utils import get_column_letter
from gspread_dataframe import get_as_dataframe

# Built-in profile defaults reproduce the Macosko-Pipelines results behavior.
# Pass --config <profile.json> to override these values (e.g. bican.json, the
# same file submit.py uses) and serve as a drop-in for another workspace.
# Unknown keys are ignored, so a single profile file can be shared between
# submit.py and results.py.
DEFAULT_PROFILE = {
    "sa": "pipelines",
    "sheet_key": "1NOaWXARQiSA6fquOtcouQPREPN4buYIf13tq_F6D9As",
    "bucket": "fc-secure-d99fbd65-eb27-4989-95b4-4cf559aa7d36",
    "fastq_match": "_S",
    "process_singlecell": True,
    "web_summary_suffix": "/outs/web_summary.html",
    "slidetags_fastq_index": "RNAIndex",
}

def get_args():
    parser = argparse.ArgumentParser(description='Terra results update script')
    parser.add_argument("--config", type=str, default=None, help="Path to JSON profile overriding the built-in Macosko defaults")
    parser.add_argument("--sa", type=str, default=None, help="Service account to impersonate (short name or full email) for Google Sheets/Drive access; overrides the profile 'sa'")
    parser.add_argument("--dryrun", action='store_true', help="Compute and print everything but do not write to the sheet")
    return parser.parse_args()

args = get_args()

# Load the profile (built-in defaults, overridden by --config)
profile = json.loads(json.dumps(DEFAULT_PROFILE))  # deep copy
if args.config is not None:
    with open(args.config) as f:
        profile.update(json.load(f))  # unknown (submit-only) keys are harmless
if args.sa is not None:
    profile["sa"] = args.sa
dryrun = args.dryrun
print(f"  config: {args.config}")
print(f"  dryrun: {dryrun}")

def retry(fn, *args, retries=5, **kwargs):
    for i in range(retries):
        try:
            return fn(*args, **kwargs)
        except gspread.exceptions.APIError as e:
            if i == retries - 1 or e.response.status_code != 500:
                raise
            time.sleep(2 ** i)

# Load the sheet using impersonated service account credentials
# (avoids needing a service account key file - uses ADC + impersonation)
source_creds, gcp_project = google_auth_default(scopes=["https://www.googleapis.com/auth/cloud-platform"])
if "@" in profile["sa"]:
    sa_email = profile["sa"]
else:
    assert gcp_project, "Could not infer project from ADC; pass a full SA email via --sa or profile 'sa'"
    sa_email = f"{profile['sa']}@{gcp_project}.iam.gserviceaccount.com"
sheets_creds = impersonated_credentials.Credentials(
    source_credentials=source_creds,
    target_principal=sa_email,
    target_scopes=[
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive",
    ],
    lifetime=3600,
)
sh = gspread.authorize(sheets_creds).open_by_key(profile["sheet_key"])

# Load the bucket
BUCKET = profile["bucket"]
bucket = storage.Client().bucket(BUCKET)
bucket.reload()

# Load the number of FASTQ files for each BCL/Index
fastqs = [blob.name for blob in bucket.list_blobs(prefix=f"fastqs") if blob.name.endswith(".fastq.gz")]
fastqs = Counter([(f.split("/")[1], f.split("/")[2].split(profile["fastq_match"])[0]) for f in fastqs])
fastqs = pd.DataFrame([(k[0],k[1],v) for k,v in fastqs.items()], columns=["BCL", "Index", "FASTQs"])
assert not fastqs.duplicated(subset=["BCL", "Index"]).any()

def blob2link(blob):
    if not isinstance(blob, str):
        return pd.NA
    elif '\n' in blob:
        return blob
    else:
        return f'=HYPERLINK("https://storage.cloud.google.com/{BUCKET}/{blob}", "{os.path.basename(blob)}")'

def paths_to_cell(paths):
    # `paths` is a list of raw blob paths sharing one BCL/Index (or NaN if unmatched).
    # A single path becomes a normal HYPERLINK() formula; multiple paths can't share one
    # formula cell, so they're written as plain text here and turned into a rich-text
    # cell with one real hyperlink per line via multiline_hyperlink_requests().
    if not isinstance(paths, list) or len(paths) == 0:
        return ''
    if len(paths) == 1:
        return blob2link(paths[0])
    return '\n'.join(paths)

def multiline_hyperlink_requests(sheet_id, col_idx, paths_series):
    requests = []
    for i, paths in enumerate(paths_series):
        if not isinstance(paths, list) or len(paths) < 2:
            continue
        runs, pos = [], 0
        for path in paths:
            runs.append({"startIndex": pos, "format": {"link": {"uri": f"https://storage.cloud.google.com/{BUCKET}/{path}"}}})
            pos += len(path) + 1
        requests.append({
            "updateCells": {
                "rows": [{"values": [{
                    "userEnteredValue": {"stringValue": "\n".join(paths)},
                    "textFormatRuns": runs,
                }]}],
                "fields": "userEnteredValue,textFormatRuns",
                "range": {
                    "sheetId": sheet_id,
                    "startRowIndex": i + 1,
                    "endRowIndex": i + 2,
                    "startColumnIndex": col_idx,
                    "endColumnIndex": col_idx + 1,
                },
            }
        })
    return requests

link_requests = []

# Query the gene-expression / slide-tags output blobs once (shared by both GEX tabs)
count_blobs = [blob.name for blob in bucket.list_blobs(prefix=f"gene-expression") if blob.name.endswith(profile["web_summary_suffix"])]
count_df = [(c.split("/")[1], c.split("/")[2], blob2link(c)) for c in count_blobs]
count_df = pd.DataFrame(count_df, columns=["BCL", "Index", "web_summary"])
assert not count_df.duplicated(subset=["BCL", "Index"]).any()

tags_blobs = [blob.name for blob in bucket.list_blobs(prefix=f"slide-tags") if blob.name.endswith("/summary.pdf")]
tags_df = [(t.split("/")[1], t.split("/")[2], t) for t in tags_blobs]
tags_df = pd.DataFrame(tags_df, columns=["BCL", "Index", "summary"])
tags_df = tags_df.groupby(['BCL', 'Index'], as_index=False).agg({'summary': list})
assert not tags_df.duplicated(subset=["BCL", "Index"]).any()

def update_gex_tab(ws_name, fastq_index_col):
    # Shared logic for the Slide-tags and SingleCell worksheets
    ws = sh.worksheet(ws_name)
    df0 = retry(get_as_dataframe, ws)
    ranges = {col: get_column_letter(df0.columns.get_loc(col)+1)+"2" for col in ["web_summary", "summary", "FASTQs"]}
    col_idx = {col: df0.columns.get_loc(col) for col in ["web_summary", "summary", "FASTQs"]}
    df0 = df0.drop(columns=["web_summary", "summary", "FASTQs"])
    df0.rename(columns={'RNAIndex': 'Index'}, inplace=True)

    df = df0.copy()
    df = df.dropna(subset=['BCL', 'Index'])
    dups = df.duplicated(subset=["BCL", "Index"])
    assert not dups.any(), f"{ws_name} sheet has duplicated BCL/RNAIndex pair:\n{df[dups]}"

    count =      pd.merge(df0, count_df, on=["BCL", "Index"], how="left").fillna('')["web_summary"]
    tags_paths = pd.merge(df0, tags_df,  on=["BCL", "Index"], how="left")["summary"]
    if fastq_index_col == "SBIndex":
        fq = pd.merge(df0, fastqs.rename(columns={"Index": "SBIndex"}), on=["BCL", "SBIndex"], how="left").fillna('')["FASTQs"]
    else:
        fq = pd.merge(df0, fastqs, on=["BCL", "Index"], how="left").fillna('')["FASTQs"]
    assert len(count) == len(tags_paths) == len(fq)

    tags = tags_paths.apply(paths_to_cell)

    if not dryrun:
        retry(ws.update, values=[[v] for v in count], range_name=ranges["web_summary"], raw=False)
        retry(ws.update, values=[[v] for v in tags],  range_name=ranges["summary"],     raw=False)
        retry(ws.update, values=[[v] for v in fq],    range_name=ranges["FASTQs"],      raw=False)

    reqs = multiline_hyperlink_requests(ws.id, col_idx["summary"], tags_paths)
    print(f"[{ws_name}] rows={len(df0)} web_summary_links={(count != '').sum()} "
          f"summary_cells={(tags != '').sum()} multiline={len(reqs)} fastq_index={fastq_index_col}")
    return reqs

################################################################################

# Load <Recon> worksheet
recon_ws = sh.worksheet("Recon")
df0 = retry(get_as_dataframe, recon_ws)
ranges = {col: get_column_letter(df0.columns.get_loc(col)+1)+"2" for col in ["QC", "summary", "FASTQs"]}
col_idx = {col: df0.columns.get_loc(col) for col in ["QC", "summary", "FASTQs"]}
df0 = df0.drop(columns=["QC", "summary", "FASTQs"])

df = df0.copy()
df = df.dropna(subset=['BCL', 'Index'])
dups = df.duplicated(subset=["BCL", "Index"])
assert not dups.any(), f"Recon sheet has duplicated BCL/Index pair:\n{df[dups]}"

# Query all PDF blobs
recon = [blob.name for blob in bucket.list_blobs(prefix=f"recon") if blob.name.endswith(".pdf")]

# Load QC.pdf blobs
QC = [(r.split("/")[1], r.split("/")[2], r) for r in recon if r.endswith("QC.pdf")]
QC = pd.DataFrame(QC, columns=["BCL", "Index", "QC"])
QC = QC.groupby(['BCL', 'Index'], as_index=False).agg({'QC': list})
assert not QC.duplicated(subset=["BCL", "Index"]).any()

# Load summary.pdf blobs
summary = [(r.split("/")[1], r.split("/")[2], r) for r in recon if r.endswith("summary.pdf")]
summary = pd.DataFrame(summary, columns=["BCL", "Index", "summary"])
summary = summary.groupby(['BCL', 'Index'], as_index=False).agg({'summary': list})
assert not summary.duplicated(subset=["BCL", "Index"]).any()

# Update the sheet
QC_paths      = pd.merge(df0, QC,      on=["BCL", "Index"], how="left")["QC"]
summary_paths = pd.merge(df0, summary, on=["BCL", "Index"], how="left")["summary"]
recon_fastqs  = pd.merge(df0, fastqs,  on=["BCL", "Index"], how="left").fillna('')["FASTQs"]
assert len(QC_paths) == len(summary_paths) == len(recon_fastqs)

QC_cells = QC_paths.apply(paths_to_cell)
summary_cells = summary_paths.apply(paths_to_cell)

if not dryrun:
    retry(recon_ws.update, values=[[v] for v in QC_cells],      range_name=ranges["QC"],      raw=False)
    retry(recon_ws.update, values=[[v] for v in summary_cells], range_name=ranges["summary"], raw=False)
    retry(recon_ws.update, values=[[v] for v in recon_fastqs],  range_name=ranges["FASTQs"],  raw=False)

link_requests += multiline_hyperlink_requests(recon_ws.id, col_idx["QC"], QC_paths)
link_requests += multiline_hyperlink_requests(recon_ws.id, col_idx["summary"], summary_paths)
print(f"[Recon] rows={len(df0)} QC_cells={(QC_cells != '').sum()} "
      f"summary_cells={(summary_cells != '').sum()} multiline={len(link_requests)}")

################################################################################

# Load <Slide-tags> worksheet
link_requests += update_gex_tab("Slide-tags", profile["slidetags_fastq_index"])

################################################################################

# Load <SingleCell> worksheet (Macosko only)
if profile["process_singlecell"]:
    link_requests += update_gex_tab("SingleCell", "RNAIndex")

################################################################################

# Apply accumulated rich-text (multiline hyperlink) cells in one batch
if link_requests and not dryrun:
    retry(sh.batch_update, {"requests": link_requests})

if dryrun:
    print(f"Dry run complete - {len(link_requests)} multiline cells would be written; no sheet updates made")
